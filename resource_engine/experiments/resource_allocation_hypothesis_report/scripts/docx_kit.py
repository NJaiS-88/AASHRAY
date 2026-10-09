"""Small python-docx toolkit for the hypothesis report: styles, tables, figures and OMML math.

Design values follow the series reference report (Word, Arial): navy table headers, blue
heading rules, light-blue hypothesis labels, restrained green/yellow conclusion cells.
"""

import re

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT = "Arial"
NAVY = "0F172A"
BLUE = "1D4ED8"
LBLUE = "DBEAFE"
ZEBRA = ("FFFFFF", "F8FAFC")
BOX = "F8FAFC"
BORDER = "CBD5E1"
GREEN = "DCFCE7"
YELLOW = "FEF9C3"
YELLOW_EDGE = "FDE68A"
BLUE_EDGE = "BFDBFE"
GREY = "64748B"
INK = "0F172A"
BODY = "111827"

BODY_PT = 10
TABLE_PT = 9
CONTENT_W = 6.87  # inches between the margins

TOKEN = re.compile(r"(\*\*|\*|\^\{[^}]*\}|_\{[^}]*\})")

# Children that follow w:pBdr inside w:pPr (borders must be inserted before them).
PPR_AFTER_PBDR = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
                  "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
                  "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
                  "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
                  "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
                  "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")


def insert_before(parent, child, successors):
    for tag in successors:
        found = parent.find(qn(tag))
        if found is not None:
            found.addprevious(child)
            return child
    parent.append(child)
    return child


# ---------------------------------------------------------------- runs and paragraphs
def style_run(run, size=None, color=None, bold=None, italic=None, font=FONT):
    run.font.name = font
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    rfonts.set(qn("w:eastAsia"), font)
    rfonts.set(qn("w:cs"), font)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    return run


def add_text(p, text, size=BODY_PT, color=BODY, bold=False, italic=False):
    """Add runs to p. Markup: **bold**, *italic*, ^{superscript}, _{subscript}."""
    b, i = bold, italic
    for part in TOKEN.split(text):
        if not part:
            continue
        if part == "**":
            b = not b
            continue
        if part == "*":
            i = not i
            continue
        sup, sub = part.startswith("^{"), part.startswith("_{")
        if sup or sub:
            part = part[2:-1]
        run = style_run(p.add_run(part), size, color, b, i)
        if sup:
            run.font.superscript = True
        if sub:
            run.font.subscript = True
    return p


def fmt(p, before=0, after=4, align=None, keep=False, line=1.1, indent=None, first=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.widow_control = True
    if keep:
        pf.keep_with_next = True
    if align is not None:
        pf.alignment = align
    if indent is not None:
        pf.left_indent = Inches(indent)
    if first is not None:
        pf.first_line_indent = Inches(first)
    return p


def bottom_rule(p, color=BLUE, size=8, space=2):
    ppr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", str(size)), ("w:space", str(space)), ("w:color", color)):
        bottom.set(qn(k), v)
    pbdr.append(bottom)
    insert_before(ppr, pbdr, PPR_AFTER_PBDR)


# ---------------------------------------------------------------- styles
def _clear_theme_fonts(rpr, color=None):
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in list(rfonts.attrib):
        del rfonts.attrib[attr]
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(k), FONT)
    if color is not None:
        old = rpr.find(qn("w:color"))
        if old is not None:
            rpr.remove(old)
        el = OxmlElement("w:color")
        el.set(qn("w:val"), color)
        # w:color sits after the font/bold/italic group; before w:sz keeps the schema order.
        insert_before(rpr, el, ("w:spacing", "w:w", "w:kern", "w:position", "w:sz", "w:szCs",
                                "w:highlight", "w:u", "w:effect", "w:bdr", "w:shd", "w:fitText",
                                "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang", "w:eastAsianLayout",
                                "w:specVanish", "w:oMath"))


def setup_styles(doc):
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    rpr = defaults.find(qn("w:rPrDefault")).find(qn("w:rPr"))
    _clear_theme_fonts(rpr)
    normal = doc.styles["Normal"]
    _clear_theme_fonts(normal.element.get_or_add_rPr(), BODY)
    normal.font.size = Pt(BODY_PT)
    for name, size, color in (("Heading 1", 15, INK), ("Heading 2", 11, "000000"), ("Heading 3", 12, BLUE)):
        st = doc.styles[name]
        _clear_theme_fonts(st.element.get_or_add_rPr(), color)
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.italic = False
        st.paragraph_format.keep_with_next = True


# ---------------------------------------------------------------- tables
def _table_props(table, widths, borders=True, border_color=BORDER, margins=(45, 45, 90, 90)):
    tbl = table._tbl
    tblpr = tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    tblw.set(qn("w:type"), "dxa")
    tblw.set(qn("w:w"), str(int(round(sum(widths) * 1440))))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    look = tblpr.find(qn("w:tblLook"))

    def put(el):
        if look is not None:
            look.addprevious(el)
        else:
            tblpr.append(el)

    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        if borders:
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "4")
            e.set(qn("w:space"), "0")
            e.set(qn("w:color"), border_color)
        else:
            e.set(qn("w:val"), "nil")
        b.append(e)
    put(b)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    put(layout)
    mar = OxmlElement("w:tblCellMar")
    for edge, val in zip(("top", "left", "bottom", "right"), (margins[0], margins[2], margins[1], margins[3])):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:w"), str(val))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    put(mar)
    grid = tbl.tblGrid
    for col, w in zip(grid.findall(qn("w:gridCol")), widths):
        col.set(qn("w:w"), str(int(round(w * 1440))))
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = Inches(w)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    va = tcpr.find(qn("w:vAlign"))
    if va is not None:
        va.addprevious(shd)
    else:
        tcpr.append(shd)


def _row_flags(row, header=False):
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))
    if header:
        trpr.append(OxmlElement("w:tblHeader"))


def cell_text(cell, text, size=TABLE_PT, color=BODY, bold=False, italic=False, align=None):
    p = cell.paragraphs[0]
    fmt(p, 0, 0, align=align, line=1.0)
    add_text(p, text, size=size, color=color, bold=bold, italic=italic)
    return p


AMAP = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT}


def data_table(doc, headers, rows, widths, align=None, fills=None, size=TABLE_PT, bold_first=True,
               bold_rows=(), keep_together=True, header_align=None):
    """Navy-header table. rows: list of lists of markup strings. fills: {(row, col): hex}."""
    fills = fills or {}
    ncol = len(headers)
    align = align or ["left"] * ncol
    header_align = header_align or align
    table = doc.add_table(rows=1 + len(rows), cols=ncol)
    _table_props(table, widths)
    hdr = table.rows[0]
    _row_flags(hdr, header=True)
    for c, text in enumerate(headers):
        cell = hdr.cells[c]
        shade(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell_text(cell, text, size=size, color="FFFFFF", bold=True, align=AMAP[header_align[c]])
    for r, values in enumerate(rows, start=1):
        row = table.rows[r]
        _row_flags(row)
        for c, text in enumerate(values):
            cell = row.cells[c]
            shade(cell, fills.get((r - 1, c), ZEBRA[(r - 1) % 2]))
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            bold = (bold_first and c == 0) or (r - 1) in bold_rows
            cell_text(cell, text, size=size, bold=bold, align=AMAP[align[c]])
    # The header row always stays with the first data row; keep_together keeps the whole table on one page.
    for row in (table.rows[:-1] if keep_together else table.rows[:1]):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = True
    return table


def box(doc, fill=BOX, edge=BORDER, width=CONTENT_W):
    """One-cell shaded box; returns the cell (its first paragraph is empty)."""
    table = doc.add_table(rows=1, cols=1)
    _table_props(table, [width], border_color=edge, margins=(80, 80, 120, 120))
    _row_flags(table.rows[0])
    cell = table.rows[0].cells[0]
    shade(cell, fill)
    fmt(cell.paragraphs[0], 0, 0, line=1.1)
    return cell


def label_box(doc, label, widths=(0.8, CONTENT_W - 0.8), keep=False):
    """Reference-style hypothesis box: light-blue label cell + light text cell. Returns the text cell.

    keep=True keeps the box on the same page as whatever follows it.
    """
    table = doc.add_table(rows=1, cols=2)
    _table_props(table, list(widths), margins=(90, 90, 110, 110))
    _row_flags(table.rows[0])
    lab, txt = table.rows[0].cells
    shade(lab, LBLUE)
    shade(txt, BOX)
    lab.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    txt.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell_text(lab, label, size=13, color=INK, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    fmt(txt.paragraphs[0], 0, 0, line=1.1)
    if keep:
        for cell in (lab, txt):
            cell.paragraphs[0].paragraph_format.keep_with_next = True
    return txt


H0_FILL = "EEF2F7"   # light grey-blue label (null hypothesis), as in the reference report
OBJ_FILL = "F1F5F9"  # neutral label for secondary evaluation objectives
SLATE = "334155"


def hypothesis_box(doc, label, label_fill, sublabel=None, label_color=INK, label_w=0.95,
                   margins=(150, 170, 150, 170)):
    """Reference-style hypothesis box: top-aligned bold label cell and a white statement cell.

    margins are cell paddings in twips: (top, bottom, left, right).
    Returns (table, text_cell); the caller fills the text cell.
    """
    table = doc.add_table(rows=1, cols=2)
    _table_props(table, [label_w, CONTENT_W - label_w], margins=margins)
    _row_flags(table.rows[0])
    lab, txt = table.rows[0].cells
    shade(lab, label_fill)
    shade(txt, "FFFFFF")
    lab.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    txt.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    cell_text(lab, label, size=13, color=label_color, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    if sublabel:
        p = lab.add_paragraph()
        fmt(p, 2, 0, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.0)
        add_text(p, sublabel, size=7.5, color=GREY)
    fmt(txt.paragraphs[0], 0, 0, line=1.15)
    return table, txt


def keep_table_with_next(table):
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = True


# ---------------------------------------------------------------- figures
def picture(p, path, width_in):
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width_in))
    return run


def native_width(path, dpi=220, max_w=6.75):
    from PIL import Image
    with Image.open(path) as im:
        return min(im.size[0] / dpi, max_w)


def caption(container, text, after=8):
    p = container.add_paragraph()
    fmt(p, 3, after, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.0)
    add_text(p, text, size=8.5, color=GREY, italic=True)
    return p


def figure(doc, path, text, width=None):
    p = doc.add_paragraph()
    fmt(p, 6, 0, align=WD_ALIGN_PARAGRAPH.CENTER, keep=True, line=1.0)
    picture(p, path, width or native_width(path))
    return caption(doc, text)


def figure_pair(doc, left, right):
    """Two half-width figures side by side, each with its own caption: (path, caption)."""
    table = doc.add_table(rows=1, cols=2)
    half = CONTENT_W / 2
    _table_props(table, [half, half], borders=False, margins=(0, 0, 40, 40))
    _row_flags(table.rows[0])
    for cell, (path, text) in zip(table.rows[0].cells, (left, right)):
        p = cell.paragraphs[0]
        fmt(p, 4, 0, align=WD_ALIGN_PARAGRAPH.CENTER, keep=True, line=1.0)
        picture(p, path, min(native_width(path), half - 0.1))
        caption(cell, text, after=2)
    return table


# ---------------------------------------------------------------- headers / footers
def add_field(p, instr, size, color):
    def fld(kind):
        r = OxmlElement("w:r")
        c = OxmlElement("w:fldChar")
        c.set(qn("w:fldCharType"), kind)
        r.append(c)
        return r

    p._p.append(fld("begin"))
    run = style_run(p.add_run(), size, color)
    t = OxmlElement("w:instrText")
    t.set(qn("xml:space"), "preserve")
    t.text = f" {instr} "
    run._r.append(t)
    p._p.append(fld("separate"))
    style_run(p.add_run("1"), size, color)
    p._p.append(fld("end"))


def footer_text(footer):
    p = footer.paragraphs[0]
    fmt(p, 0, 0, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.0)
    add_text(p, "AASHRAY | Hypothesis Testing Report | Page ", size=8, color=GREY)
    add_field(p, "PAGE", 8, GREY)


def header_text(header, right_text):
    p = header.paragraphs[0]
    fmt(p, 0, 0, line=1.0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(CONTENT_W), WD_TAB_ALIGNMENT.RIGHT)
    add_text(p, "AASHRAY", size=8, color=BLUE, bold=True)
    add_text(p, "\t" + right_text, size=8, color=GREY)
    bottom_rule(p, color=BORDER, size=4, space=3)


# ---------------------------------------------------------------- OMML math
M_NS = ('xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" '
        'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"')


class Math:
    """Builds OMML fragments; size is in points."""

    def __init__(self, size=10.5):
        self.sz = int(round(size * 2))

    @staticmethod
    def _esc(t):
        return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    def r(self, text, plain=False):
        sty = '<m:rPr><m:sty m:val="p"/></m:rPr>' if plain else ""
        return (f'<m:r>{sty}<w:rPr><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math"/>'
                f'<w:sz w:val="{self.sz}"/></w:rPr><m:t xml:space="preserve">{self._esc(text)}</m:t></m:r>')

    def t(self, text):
        return self.r(text, plain=True)

    def words(self, text):
        """Normal (non-math) text inside an equation: hyphens stay hyphens, not minus signs."""
        return (f'<m:r><m:rPr><m:nor/></m:rPr><w:rPr><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math"/>'
                f'<w:sz w:val="{self.sz}"/></w:rPr><m:t xml:space="preserve">{self._esc(text)}</m:t></m:r>')

    @staticmethod
    def sub(base, s):
        return f"<m:sSub><m:e>{base}</m:e><m:sub>{s}</m:sub></m:sSub>"

    @staticmethod
    def sup(base, s):
        return f"<m:sSup><m:e>{base}</m:e><m:sup>{s}</m:sup></m:sSup>"

    @staticmethod
    def subsup(base, s, p):
        return f"<m:sSubSup><m:e>{base}</m:e><m:sub>{s}</m:sub><m:sup>{p}</m:sup></m:sSubSup>"

    @staticmethod
    def frac(num, den):
        return f"<m:f><m:num>{num}</m:num><m:den>{den}</m:den></m:f>"

    @staticmethod
    def nary(ch, sub, body, sup="", loc="undOvr"):
        hide = '<m:supHide m:val="1"/>' if not sup else ""
        return (f'<m:nary><m:naryPr><m:chr m:val="{ch}"/><m:limLoc m:val="{loc}"/>{hide}</m:naryPr>'
                f"<m:sub>{sub}</m:sub><m:sup>{sup}</m:sup><m:e>{body}</m:e></m:nary>")

    @staticmethod
    def delim(inner, beg="(", end=")"):
        return (f'<m:d><m:dPr><m:begChr m:val="{beg}"/><m:endChr m:val="{end}"/></m:dPr>'
                f"<m:e>{inner}</m:e></m:d>")

    @staticmethod
    def eqarr(*rows):
        return "<m:eqArr>" + "".join(f"<m:e>{x}</m:e>" for x in rows) + "</m:eqArr>"

    @staticmethod
    def matrix(rows, jc="left", gap_twips=None):
        """Matrix with left-justified columns; used for case definitions and constraint lists."""
        cols = max(len(r) for r in rows)
        gap = f'<m:cGpRule m:val="3"/><m:cGp m:val="{gap_twips}"/>' if gap_twips else ""
        mpr = (f'<m:mPr><m:baseJc m:val="center"/>{gap}<m:mcs><m:mc><m:mcPr><m:count m:val="{cols}"/>'
               f'<m:mcJc m:val="{jc}"/></m:mcPr></m:mc></m:mcs></m:mPr>')
        body = "".join("<m:mr>" + "".join(f"<m:e>{c}</m:e>" for c in row) + "</m:mr>" for row in rows)
        return f"<m:m>{mpr}{body}</m:m>"

    @staticmethod
    def bar(inner):
        return f'<m:bar><m:barPr><m:pos m:val="top"/></m:barPr><m:e>{inner}</m:e></m:bar>'

    @staticmethod
    def func(name_xml, arg):
        return f"<m:func><m:fName>{name_xml}</m:fName><m:e>{arg}</m:e></m:func>"

    @staticmethod
    def inline(inner):
        return parse_xml(f"<m:oMath {M_NS}>{inner}</m:oMath>")

    @staticmethod
    def display(inner, jc="center"):
        return parse_xml(f'<m:oMathPara {M_NS}><m:oMathParaPr><m:jc m:val="{jc}"/></m:oMathParaPr>'
                         f"<m:oMath>{inner}</m:oMath></m:oMathPara>")


def math_para(container, xml, before=2, after=4, keep=False):
    p = container.add_paragraph()
    fmt(p, before, after, align=WD_ALIGN_PARAGRAPH.CENTER, keep=keep, line=1.0)
    p._p.append(Math.display(xml))
    return p


def bullet(container, text, symbol="•", indent=0.28, size=BODY_PT, after=3):
    p = container.add_paragraph()
    fmt(p, 0, after, indent=indent, first=-0.18, line=1.1)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(indent))
    add_text(p, f"{symbol}\t{text}", size=size)
    return p
